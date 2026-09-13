"""generate_defense_ppt — 生成智慧城市比赛答辩 PPTX。

所属层：scripts
依赖：python-pptx, Pillow
对接算法层：N/A
"""
from pathlib import Path
from typing import Iterable, Optional

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "deliverables" / "assets"
OUTPUT = ROOT / "deliverables" / "TrafficGraph_智慧城市答辩.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

COLORS = {
    "ink": RGBColor(15, 30, 24),
    "ink_soft": RGBColor(42, 61, 51),
    "paper": RGBColor(246, 249, 247),
    "white": RGBColor(255, 255, 255),
    "line": RGBColor(211, 221, 215),
    "muted": RGBColor(91, 108, 99),
    "teal": RGBColor(15, 145, 100),
    "teal_soft": RGBColor(229, 245, 238),
    "amber": RGBColor(213, 145, 31),
    "amber_soft": RGBColor(255, 246, 224),
    "red": RGBColor(190, 64, 52),
    "red_soft": RGBColor(253, 238, 235),
    "dark_card": RGBColor(20, 39, 31),
}


def set_background(slide, color: RGBColor) -> None:
    """设置幻灯片纯色背景。"""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 18,
    color: RGBColor = COLORS["ink"],
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    font: str = "Microsoft YaHei",
) -> None:
    """添加统一字体和颜色的文本框。"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bullets(
    slide,
    items: Iterable[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 17,
    color: RGBColor = COLORS["ink_soft"],
    bullet_color: RGBColor = COLORS["teal"],
) -> None:
    """添加项目符号列表。"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(7)
        marker = paragraph.add_run()
        marker.text = "● "
        marker.font.name = "Microsoft YaHei"
        marker.font.size = Pt(max(8, size - 5))
        marker.font.color.rgb = bullet_color
        run = paragraph.add_run()
        run.text = item
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.color.rgb = color


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    body: str,
    *,
    tone: str = "teal",
) -> None:
    """添加信息卡片。"""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS["white"]
    card.line.color.rgb = COLORS["line"]
    card.line.width = Pt(0.8)
    card.adjustments[0] = 0.05

    accent_color = COLORS[tone]
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(0.06),
        Inches(h),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = accent_color
    accent.line.fill.background()

    add_text(slide, title, x + 0.22, y + 0.18, w - 0.4, 0.35, size=17, bold=True)
    add_text(
        slide,
        body,
        x + 0.22,
        y + 0.7,
        w - 0.4,
        h - 0.85,
        size=13,
        color=COLORS["muted"],
    )


def add_header(slide, title: str, subtitle: str = "", page: Optional[int] = None) -> None:
    """添加内容页页眉。"""
    add_text(slide, title, 0.65, 0.35, 11.8, 0.5, size=25, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.67, 0.88, 11.8, 0.35, size=12, color=COLORS["muted"])
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.65),
        Inches(1.23),
        Inches(12.0),
        Inches(0.018),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.fill.background()
    if page is not None:
        add_text(
            slide,
            f"{page:02d}",
            12.32,
            6.94,
            0.5,
            0.25,
            size=10,
            color=COLORS["muted"],
            align=PP_ALIGN.RIGHT,
        )


def add_image_contain(slide, path: Path, x: float, y: float, w: float, h: float) -> None:
    """按比例把图片放入指定区域。"""
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = w / h
    if image_ratio > box_ratio:
        image_w = w
        image_h = w / image_ratio
    else:
        image_h = h
        image_w = h * image_ratio
    left = x + (w - image_w) / 2
    top = y + (h - image_h) / 2
    slide.shapes.add_picture(
        str(path),
        Inches(left),
        Inches(top),
        width=Inches(image_w),
        height=Inches(image_h),
    )


def add_architecture(slide, x: float, y: float, w: float) -> None:
    """绘制五层架构图。"""
    layers = [
        ("应用层", "智慧交通指挥 Web 控制台", "teal"),
        ("Agent 决策层", "认知解析 · 工具调度 · 报告生成", "amber"),
        ("业务 Skill 层", "交通专家 · 信号调度 · UI Router", "teal"),
        ("工具与数据层", "12 个交通数据工具 · TrafficSim", "amber"),
        ("知识基础设施", "ChromaDB RAG · 长期记忆 · 审计", "teal"),
    ]
    layer_h = 0.78
    gap = 0.12
    for index, (title, body, tone) in enumerate(layers):
        top = y + index * (layer_h + gap)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(top),
            Inches(w),
            Inches(layer_h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS["dark_card"] if index in (1, 3) else COLORS["white"]
        shape.line.color.rgb = COLORS["line"]
        shape.adjustments[0] = 0.04
        add_text(
            slide,
            title,
            x + 0.18,
            top + 0.12,
            1.5,
            0.3,
            size=14,
            bold=True,
            color=COLORS["white"] if index in (1, 3) else COLORS["ink"],
        )
        add_text(
            slide,
            body,
            x + 1.7,
            top + 0.13,
            w - 1.9,
            0.35,
            size=12,
            color=RGBColor(218, 231, 224) if index in (1, 3) else COLORS["muted"],
        )


def create_deck() -> Path:
    """创建答辩 PPTX 文件。"""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    # 1. 封面
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["ink"])
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.72),
        Inches(1.18),
        Inches(0.09),
        Inches(2.16),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["teal"]
    bar.line.fill.background()
    add_text(slide, "TrafficGraph", 1.05, 1.12, 9.0, 0.72, size=38, bold=True, color=COLORS["white"])
    add_text(
        slide,
        "上海市智慧交通指挥 Agent",
        1.05,
        1.95,
        10.6,
        0.62,
        size=28,
        color=RGBColor(211, 230, 220),
    )
    add_text(
        slide,
        "多源态势融合 · 信号优化 · 事件处置 · 数据可视化 · 智能报表",
        1.07,
        2.82,
        11.3,
        0.42,
        size=15,
        color=RGBColor(148, 180, 164),
    )
    add_text(
        slide,
        "智慧城市竞赛答辩",
        1.07,
        5.9,
        3.0,
        0.35,
        size=14,
        color=COLORS["white"],
    )
    add_text(
        slide,
        "2026",
        10.8,
        5.9,
        1.4,
        0.35,
        size=14,
        color=RGBColor(137, 171, 154),
        align=PP_ALIGN.RIGHT,
    )

    # 2. 背景与痛点
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "项目背景与问题", "交通指挥不缺少数据，缺少把数据快速变成处置判断的能力", 2)
    add_card(slide, 0.75, 1.62, 3.75, 2.25, "数据分散", "路网、信号、事件、警情、视频和数据报表分散在不同系统，值班人员需要反复切换查询。", tone="red")
    add_card(slide, 4.78, 1.62, 3.75, 2.25, "研判依赖经验", "高峰拥堵原因复杂，传统监控只能看到结果，难以快速关联饱和度、排队、配时和事件影响。", tone="amber")
    add_card(slide, 8.81, 1.62, 3.75, 2.25, "决策链路长", "从发现问题到形成处置建议需要人工汇总，缺少统一的分析、解释和报告输出链路。", tone="teal")
    add_text(slide, "项目目标", 0.75, 4.35, 2.0, 0.35, size=20, bold=True)
    add_bullets(
        slide,
        [
            "让自然语言成为交通指挥系统的统一查询入口。",
            "把实时数据、专业知识、信号优化和事件处置串联到一个 Agent。",
            "用一屏式控制台完成态势、诊断、建议、导出和审计。",
        ],
        0.8,
        4.88,
        11.7,
        1.7,
        size=16,
    )

    # 3. 产品定位
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "产品定位", "面向交通指挥场景的决策辅助 Agent，而不是自动控制系统", 3)
    add_card(slide, 0.75, 1.6, 5.8, 2.15, "TrafficGraph 做什么", "理解用户意图，调用交通数据工具和专业 RAG，对全市、路口、干线、信号和事件进行分析，并生成可追溯的研判报告。", tone="teal")
    add_card(slide, 6.78, 1.6, 5.8, 2.15, "TrafficGraph 不做什么", "不直接修改真实信号配时，不发布交通管制，不代替值班长审批。所有行动建议都保留人工确认环节。", tone="red")
    add_text(slide, "四个核心工作视图", 0.75, 4.12, 3.0, 0.4, size=20, bold=True)
    view_titles = ["态势总览", "路网监控", "信号控制", "事件处置"]
    for index, title in enumerate(view_titles):
        x = 0.75 + index * 3.08
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(4.73),
            Inches(2.72),
            Inches(1.25),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS["dark_card"] if index == 0 else COLORS["white"]
        shape.line.color.rgb = COLORS["line"]
        add_text(
            slide,
            title,
            x + 0.18,
            5.08,
            2.36,
            0.35,
            size=16,
            bold=True,
            color=COLORS["white"] if index == 0 else COLORS["ink"],
            align=PP_ALIGN.CENTER,
        )

    # 4. 系统架构
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "系统架构", "从自然语言请求到交通数据、工具调用、知识检索和报表输出", 4)
    add_architecture(slide, 0.9, 1.55, 6.0)
    add_text(slide, "技术栈", 7.55, 1.55, 2.0, 0.4, size=20, bold=True)
    techs = [
        "LangGraph：Agent 状态图和工具循环",
        "FastAPI + Uvicorn：API 与 SSE 流式服务",
        "Pydantic：工具输入输出与强类型契约",
        "ChromaDB：交通专业知识库",
        "TrafficSim：确定性交通仿真",
        "原生 Web + SVG：一屏式中控台",
    ]
    add_bullets(slide, techs, 7.6, 2.18, 4.9, 3.7, size=15)

    # 5. Agent 工作流
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "Agent 工作流", "认知解析、工具调度、技能编排、报告生成和动作确认", 5)
    steps = [
        ("01", "意图理解", "识别数据查询、专业问答、页面查看或报表导出意图"),
        ("02", "工具选择", "从 12 个交通数据工具和知识工具中选择最少必要调用"),
        ("03", "技能编排", "交通专家、信号调度和 UI Router 统一处理结果"),
        ("04", "研判报告", "结合实时数据、知识来源和安全约束生成 Markdown"),
        ("05", "人工确认", "页面动作只提示，不自动覆盖当前指挥视图"),
    ]
    for index, (number, title, body) in enumerate(steps):
        x = 0.75 + index * 2.5
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(2.02),
            Inches(2.15),
            Inches(3.15),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLORS["white"]
        card.line.color.rgb = COLORS["line"]
        add_text(slide, number, x + 0.18, 2.25, 0.6, 0.35, size=13, bold=True, color=COLORS["teal"])
        add_text(slide, title, x + 0.18, 2.78, 1.8, 0.35, size=17, bold=True)
        add_text(slide, body, x + 0.18, 3.32, 1.78, 1.45, size=12, color=COLORS["muted"])
        if index < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON,
                Inches(x + 2.16),
                Inches(3.31),
                Inches(0.28),
                Inches(0.34),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COLORS["teal"]
            arrow.line.fill.background()

    # 6. 交通数据工具
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "12 个交通数据工具", "将交通业务能力标准化为 Agent 可调度的工具契约", 6)
    tools = [
        ("全市态势", "拥堵指数、车速、Top5、事件数"),
        ("路口状态", "LOS、饱和度、排队、延误、相位"),
        ("干线绿波", "车速、带宽、相位差、协调状态"),
        ("辖区指数", "排名、趋势、最堵路口"),
        ("信号配时", "周期、相位绿灯、黄灯、offset"),
        ("信号优化", "Webster 周期、绿信比、安全约束"),
        ("流量历史", "路口/干线逐时流量与车速"),
        ("活跃事件", "事故、抛锚、施工、管制"),
        ("历史事件", "处置时长、派单号、恢复状态"),
        ("视频监控", "摄像头清单、状态、预置位"),
        ("派单调度", "警力、拖车、到场时间、结果"),
        ("接处警", "110、122、热线、视频识别记录"),
    ]
    for index, (title, body) in enumerate(tools):
        column = index % 2
        row = index // 2
        x = 0.8 + column * 6.1
        y = 1.55 + row * 0.87
        add_card(slide, x, y, 5.75, 0.7, title, body, tone="teal" if index % 3 else "amber")

    # 7. RAG 与记忆
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "交通专业知识与记忆", "让回答有依据，让对话有上下文", 7)
    add_card(slide, 0.75, 1.65, 3.75, 4.25, "L1 会话记忆", "LangGraph checkpoint 按 thread_id 隔离，支持多轮对话恢复。", tone="teal")
    add_card(slide, 4.78, 1.65, 3.75, 4.25, "L2 长期记忆", "按环境、区域、Agent 和记忆范围隔离，临时状态强制设置 TTL。", tone="amber")
    add_card(slide, 8.81, 1.65, 3.75, 4.25, "L3 知识检索", "ChromaDB 存储交通规范、信号控制、绿波、事件处置和调度流程。支持离线 hashing，也支持 BGE。", tone="teal")
    add_text(slide, "低置信度不猜测", 1.0, 5.25, 2.2, 0.35, size=16, bold=True)
    add_text(slide, "检索置信度不足时触发拒答，并提示用户补充对象、时间或规范方向。", 1.0, 5.7, 10.8, 0.5, size=14, color=COLORS["muted"])

    # 8. 中控台
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "一屏式智慧交通中控台", "四个独立工作视图，点击切换，不依赖长页面滚动", 8)
    screenshot = ASSETS / "overview.png"
    if screenshot.exists():
        add_image_contain(slide, screenshot, 0.55, 1.48, 8.35, 5.15)
    add_bullets(
        slide,
        [
            "左侧固定导航，点击切换态势、路网、信号和事件。",
            "中间工作区独立布局，右侧 Agent 固定，不带动主页面跳转。",
            "路网节点支持路口探查和诊断。",
            "浅色高对比数据面板保证投屏可读性。",
        ],
        9.15,
        1.7,
        3.55,
        4.5,
        size=14,
    )

    # 9. 信号与事件
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "信号控制与事件处置", "从实时运行状态到优化建议和处置时间线", 9)
    signal_shot = ASSETS / "signals.png"
    incident_shot = ASSETS / "incidents.png"
    if signal_shot.exists():
        add_image_contain(slide, signal_shot, 0.55, 1.5, 6.05, 4.85)
    if incident_shot.exists():
        add_image_contain(slide, incident_shot, 6.82, 1.5, 6.0, 4.85)
    add_text(slide, "信号控制", 0.7, 6.55, 2.0, 0.3, size=16, bold=True, color=COLORS["teal"])
    add_text(slide, "事件处置", 6.95, 6.55, 2.0, 0.3, size=16, bold=True, color=COLORS["amber"])

    # 10. 报表与可视化
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "数据可视化与智能报表", "同一份数据同时生成表格、图表和 CSV", 10)
    report_shot = ASSETS / "report.png"
    if report_shot.exists():
        add_image_contain(slide, report_shot, 0.55, 1.5, 8.35, 5.2)
    add_bullets(
        slide,
        [
            "支持柱状图、折线图和环形图。",
            "当前工作区可一键导出报告。",
            "CSV 使用 UTF-8 BOM，可直接用 Excel 打开。",
            "Agent 返回的数据卡片也可在对话中直接渲染图表。",
        ],
        9.15,
        1.75,
        3.5,
        4.3,
        size=14,
    )

    # 11. 工程与安全
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["paper"])
    add_header(slide, "工程实现与安全边界", "可演示、可复现、可扩展、不可越权", 11)
    add_card(slide, 0.75, 1.6, 3.75, 2.0, "确定性仿真", "相同日期、对象和数据类型产生稳定结果，适合现场演示和回归测试。", tone="teal")
    add_card(slide, 4.78, 1.6, 3.75, 2.0, "流式交互", "SSE 推送 thinking、tool_call、tool_result、rag_sources、text、action 和 data_card。", tone="amber")
    add_card(slide, 8.81, 1.6, 3.75, 2.0, "强类型契约", "Pydantic 定义交通数据、工具输出、UI 动作和记忆元数据。", tone="teal")
    add_text(slide, "安全边界", 0.75, 4.1, 2.0, 0.4, size=20, bold=True)
    add_bullets(
        slide,
        [
            "Agent 只能查询、分析和建议，不能直接控制真实信号。",
            "配时建议必须校验最小绿灯、行人过街和协调同步。",
            "页面动作需要人工确认，不自动执行。",
            "本地密钥只放在被忽略的 .env 文件中。",
        ],
        0.8,
        4.62,
        11.8,
        1.9,
        size=15,
        bullet_color=COLORS["red"],
    )

    # 12. 比赛亮点总结
    slide = prs.slides.add_slide(blank)
    set_background(slide, COLORS["ink"])
    add_text(slide, "项目亮点", 0.8, 0.65, 5.0, 0.55, size=30, bold=True, color=COLORS["white"])
    highlights = [
        ("1", "真实业务链", "从态势、路口、信号到事件和调度形成完整闭环。"),
        ("2", "多工具 Agent", "12 个交通数据工具统一注册、统一调度、统一返回。"),
        ("3", "知识与数据结合", "实时仿真数据与交通规范 RAG 联合研判。"),
        ("4", "一屏指挥体验", "四个固定视图、高对比中控布局和独立 Agent 面板。"),
        ("5", "自动报表输出", "表格、图表和 CSV 一键生成，支持比赛现场展示。"),
    ]
    for index, (number, title, body) in enumerate(highlights):
        y = 1.52 + index * 0.96
        number_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.85),
            Inches(y),
            Inches(0.52),
            Inches(0.52),
        )
        number_box.fill.solid()
        number_box.fill.fore_color.rgb = COLORS["teal"]
        number_box.line.fill.background()
        add_text(slide, number, 0.85, y + 0.07, 0.52, 0.3, size=15, bold=True, color=COLORS["white"], align=PP_ALIGN.CENTER)
        add_text(slide, title, 1.62, y - 0.02, 2.4, 0.34, size=17, bold=True, color=COLORS["white"])
        add_text(slide, body, 4.0, y - 0.01, 8.2, 0.42, size=14, color=RGBColor(188, 211, 198))
    add_text(slide, "谢谢评审", 9.8, 6.6, 2.8, 0.42, size=18, bold=True, color=COLORS["teal"], align=PP_ALIGN.RIGHT)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    result = create_deck()
    print(f"PPTX generated: {result}")
